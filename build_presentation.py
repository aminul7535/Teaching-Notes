#!/usr/bin/env python3
"""
Generate a professional, kid-friendly PowerPoint presentation for a
Class VI Computer classroom: Chapter 3 - More on Excel.

All content is preserved EXACTLY as provided. This script only handles the
visual formatting/layout. Every object stays editable (native text boxes and
shapes; no rasterized text).
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from pptx.enum.text import MSO_AUTO_SIZE
import copy

# ---------------------------------------------------------------------------
# Palette
# ---------------------------------------------------------------------------
DARK_BLUE   = RGBColor(0x1F, 0x38, 0x64)   # questions
BLUE        = RGBColor(0x2E, 0x75, 0xB6)   # accent blue
BLUE_DEEP   = RGBColor(0x1B, 0x5E, 0xA8)
GREEN       = RGBColor(0x2E, 0x7D, 0x32)   # answers
GREEN_DARK  = RGBColor(0x1B, 0x5E, 0x20)
GREEN_BG    = RGBColor(0xE7, 0xF6, 0xE9)   # answer box fill
GREEN_LINE  = RGBColor(0x4C, 0xAF, 0x50)
ORANGE      = RGBColor(0xED, 0x7D, 0x31)   # accent orange
ORANGE_BG   = RGBColor(0xFD, 0xEF, 0xE3)
RED         = RGBColor(0xD3, 0x2F, 0x2F)   # false icon
LIGHT_BLUE  = RGBColor(0xEA, 0xF2, 0xFB)   # question box fill
LIGHT_BLUE2 = RGBColor(0xDD, 0xEB, 0xF7)
WHITE       = RGBColor(0xFF, 0xFF, 0xFF)
TEXT_DARK   = RGBColor(0x2B, 0x2B, 0x2B)
GREY        = RGBColor(0x70, 0x70, 0x70)
BG_TOP      = RGBColor(0xFF, 0xFF, 0xFF)
BG_BOT      = RGBColor(0xEB, 0xF3, 0xFB)

TITLE_FONT = "Calibri Light"
BODY_FONT  = "Calibri"

# ---------------------------------------------------------------------------
# Presentation setup (16:9 widescreen)
# ---------------------------------------------------------------------------
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
BLANK = prs.slide_layouts[6]
SW = prs.slide_width
SH = prs.slide_height


# ---------------------------------------------------------------------------
# Low level helpers
# ---------------------------------------------------------------------------
def _set_fill(shape, color):
    shape.fill.solid()
    shape.fill.fore_color.rgb = color


def _no_line(shape):
    shape.line.fill.background()


def _line(shape, color, width_pt=1.5):
    shape.line.color.rgb = color
    shape.line.width = Pt(width_pt)


def add_soft_shadow(shape, blur=100000, dist=45000, direction=5400000, alpha=32000):
    """Add a soft outer drop shadow via raw XML."""
    spPr = shape._element.spPr
    for tag in ("a:effectLst",):
        for e in spPr.findall(qn(tag)):
            spPr.remove(e)
    effectLst = spPr.makeelement(qn("a:effectLst"), {})
    shdw = effectLst.makeelement(qn("a:outerShdw"), {
        "blurRad": str(blur),
        "dist": str(dist),
        "dir": str(direction),
        "rotWithShape": "0",
    })
    clr = shdw.makeelement(qn("a:srgbClr"), {"val": "3A4A5A"})
    a = clr.makeelement(qn("a:alpha"), {"val": str(alpha)})
    clr.append(a)
    shdw.append(clr)
    effectLst.append(shdw)
    spPr.append(effectLst)


def set_gradient_bg(slide, top=BG_TOP, bottom=BG_BOT):
    """Full-slide soft vertical gradient rectangle sent to the back."""
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    rect.shadow.inherit = False
    _no_line(rect)
    rect.fill.gradient()
    stops = rect.fill.gradient_stops
    stops[0].position = 0.0
    stops[0].color.rgb = top
    stops[1].position = 1.0
    stops[1].color.rgb = bottom
    try:
        rect.fill.gradient_angle = 90.0
    except Exception:
        pass
    # send to back
    sp = rect._element
    sp.getparent().remove(sp)
    slide.shapes._spTree.insert(2, sp)
    return rect


def _para(tf, text, size, color, bold=False, italic=False, font=BODY_FONT,
          align=PP_ALIGN.LEFT, space_before=4, space_after=4, level=0,
          line_spacing=1.05, first=False):
    p = tf.paragraphs[0] if first and not tf.paragraphs[0].runs else tf.add_paragraph()
    p.alignment = align
    p.level = level
    if space_before is not None:
        p.space_before = Pt(space_before)
    if space_after is not None:
        p.space_after = Pt(space_after)
    if line_spacing is not None:
        p.line_spacing = line_spacing
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.name = font
    r.font.color.rgb = color
    return p, r


def add_textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP, wrap=True):
    tb = slide.shapes.add_textbox(x, y, w, h)
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.vertical_anchor = anchor
    tf.margin_left = Pt(6)
    tf.margin_right = Pt(6)
    tf.margin_top = Pt(4)
    tf.margin_bottom = Pt(4)
    tf.paragraphs[0].text = ""
    return tb, tf


def add_round_rect(slide, x, y, w, h, fill, line_color=None, line_w=1.5,
                   radius=0.12, shadow=True):
    shp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h)
    shp.shadow.inherit = False
    _set_fill(shp, fill)
    if line_color is not None:
        _line(shp, line_color, line_w)
    else:
        _no_line(shp)
    try:
        shp.adjustments[0] = radius
    except Exception:
        pass
    if shadow:
        add_soft_shadow(shp)
    return shp


# ---------------------------------------------------------------------------
# Simple fade-in (entrance) animations, applied on click, in order.
# ---------------------------------------------------------------------------
def add_fade_animations(slide, shape_ids):
    if not shape_ids:
        return
    nsp = 'xmlns:p="http://schemas.openxmlformats.org/presentationml/2006/main"'
    nsa = 'xmlns:a="http://schemas.openxmlformats.org/drawingml/2006/main"'
    cid = [2]  # running id counter (id 1 is tmRoot)

    def nid():
        cid[0] += 1
        return cid[0]

    effects = []
    for sid in shape_ids:
        i1, i2, i3, i4, i5 = nid(), nid(), nid(), nid(), nid()
        effects.append(f"""
        <p:par>
          <p:cTn id="{i1}" fill="hold">
            <p:stCondLst><p:cond delay="indefinite"/></p:stCondLst>
            <p:childTnLst>
              <p:par>
                <p:cTn id="{i2}" fill="hold">
                  <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                  <p:childTnLst>
                    <p:par>
                      <p:cTn id="{i3}" presetID="10" presetClass="entr" presetSubtype="0" fill="hold" grpId="0" nodeType="clickEffect">
                        <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                        <p:childTnLst>
                          <p:animEffect transition="in" filter="fade">
                            <p:cBhvr>
                              <p:cTn id="{i4}" dur="500"/>
                              <p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl>
                            </p:cBhvr>
                          </p:animEffect>
                          <p:set>
                            <p:cBhvr>
                              <p:cTn id="{i5}" dur="1" fill="hold">
                                <p:stCondLst><p:cond delay="0"/></p:stCondLst>
                              </p:cTn>
                              <p:tgtEl><p:spTgt spid="{sid}"/></p:tgtEl>
                              <p:attrNameLst><p:attrName>style.visibility</p:attrName></p:attrNameLst>
                            </p:cBhvr>
                            <p:to><p:strVal val="visible"/></p:to>
                          </p:set>
                        </p:childTnLst>
                      </p:cTn>
                    </p:par>
                  </p:childTnLst>
                </p:cTn>
              </p:par>
            </p:childTnLst>
          </p:cTn>
        </p:par>""")

    timing_xml = f"""<p:timing {nsp} {nsa}>
      <p:tnLst>
        <p:par>
          <p:cTn id="1" dur="indefinite" restart="never" nodeType="tmRoot">
            <p:childTnLst>
              <p:seq concurrent="1" nextAc="seek">
                <p:cTn id="2" dur="indefinite" nodeType="mainSeq">
                  <p:childTnLst>{''.join(effects)}
                  </p:childTnLst>
                </p:cTn>
                <p:prevCondLst><p:cond evt="onPrev" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:prevCondLst>
                <p:nextCondLst><p:cond evt="onNext" delay="0"><p:tgtEl><p:sldTgt/></p:tgtEl></p:cond></p:nextCondLst>
              </p:seq>
            </p:childTnLst>
          </p:cTn>
        </p:par>
      </p:tnLst>
    </p:timing>"""

    from pptx.oxml import parse_xml
    timing = parse_xml(timing_xml)
    slide._element.append(timing)


# ---------------------------------------------------------------------------
# Footer + page number (added to every slide)
# ---------------------------------------------------------------------------
PAGE = [0]


def add_footer(slide, dark=False):
    PAGE[0] += 1
    n = PAGE[0]
    # thin footer accent bar
    bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, SH - Inches(0.42),
                                 SW, Inches(0.42))
    bar.shadow.inherit = False
    _no_line(bar)
    bar.fill.gradient()
    st = bar.fill.gradient_stops
    st[0].position = 0.0
    st[0].color.rgb = BLUE
    st[1].position = 1.0
    st[1].color.rgb = GREEN
    try:
        bar.fill.gradient_angle = 0.0
    except Exception:
        pass

    tb, tf = add_textbox(slide, Inches(0.35), SH - Inches(0.42),
                         Inches(8.5), Inches(0.42), anchor=MSO_ANCHOR.MIDDLE)
    _para(tf, "Chapter 3 \u2022 More on Excel   |  ", 11, WHITE, bold=True,
          first=True, space_before=0, space_after=0)

    tb2, tf2 = add_textbox(slide, SW - Inches(2.4), SH - Inches(0.42),
                           Inches(2.05), Inches(0.42), anchor=MSO_ANCHOR.MIDDLE)
    p, _ = _para(tf2, f"Page {n}", 11, WHITE, bold=True, align=PP_ALIGN.RIGHT,
                 first=True, space_before=0, space_after=0)


# ---------------------------------------------------------------------------
# Decorative scattered icon helper
# ---------------------------------------------------------------------------
def deco_icon(slide, emoji, x, y, size=28, color=None):
    tb, tf = add_textbox(slide, x, y, Inches(0.9), Inches(0.9),
                         anchor=MSO_ANCHOR.MIDDLE)
    p, r = _para(tf, emoji, size, color or WHITE, first=True,
                 align=PP_ALIGN.CENTER, space_before=0, space_after=0)
    return tb


# ---------------------------------------------------------------------------
# Slide builders
# ---------------------------------------------------------------------------
def new_slide():
    return prs.slides.add_slide(BLANK)


def content_header(slide, icon, heading, sub=None, accent=BLUE):
    """Rounded header bar at the top of a content slide."""
    bar = add_round_rect(slide, Inches(0.45), Inches(0.35),
                         SW - Inches(0.9), Inches(1.02), accent,
                         radius=0.28, shadow=True)
    tb, tf = add_textbox(slide, Inches(0.75), Inches(0.35),
                         SW - Inches(1.5), Inches(1.02),
                         anchor=MSO_ANCHOR.MIDDLE)
    p, r = _para(tf, f"{icon}  {heading}", 27, WHITE, bold=True,
                 font=TITLE_FONT, first=True, space_before=0, space_after=0)
    if sub:
        _para(tf, sub, 15, RGBColor(0xE6, 0xEF, 0xF7), italic=True,
              space_before=2, space_after=0)
    return bar


# ---- Title slide ----------------------------------------------------------
def build_title():
    slide = new_slide()
    set_gradient_bg(slide, WHITE, RGBColor(0xDF, 0xEC, 0xF9))

    # decorative spreadsheet grid panel on the right
    grid_x = Inches(8.9)
    grid_y = Inches(1.55)
    panel = add_round_rect(slide, grid_x - Inches(0.25), grid_y - Inches(0.3),
                           Inches(3.7), Inches(3.6), WHITE, radius=0.06)
    # header row of the mini spreadsheet
    cols = 4
    rows = 5
    cw = Inches(0.8)
    ch = Inches(0.55)
    cell_colors = [BLUE, GREEN, ORANGE]
    for rr in range(rows):
        for cc in range(cols):
            cx = grid_x + cc * cw
            cy = grid_y + rr * ch
            cell = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx, cy, cw, ch)
            cell.shadow.inherit = False
            _line(cell, RGBColor(0xC9, 0xD6, 0xE5), 0.75)
            if rr == 0:
                _set_fill(cell, BLUE_DEEP)
            elif cc == 0:
                _set_fill(cell, LIGHT_BLUE2)
            else:
                _set_fill(cell, WHITE)
    # a small chart emoji on grid
    deco_icon(slide, "\U0001F4CA", grid_x + Inches(2.35), grid_y - Inches(1.15),
              size=40, color=GREEN)
    deco_icon(slide, "\U0001F9EE", grid_x - Inches(0.9), grid_y + Inches(2.7),
              size=34, color=ORANGE)

    # Chapter pill
    pill = add_round_rect(slide, Inches(0.9), Inches(1.5), Inches(2.7),
                          Inches(0.75), ORANGE, radius=0.5, shadow=True)
    tb, tf = add_textbox(slide, Inches(0.9), Inches(1.5), Inches(2.7),
                         Inches(0.75), anchor=MSO_ANCHOR.MIDDLE)
    _para(tf, "\U0001F4CA  CHAPTER 3", 20, WHITE, bold=True,
          align=PP_ALIGN.CENTER, first=True, space_before=0, space_after=0)

    # Main title
    tb, tf = add_textbox(slide, Inches(0.85), Inches(2.55), Inches(7.6),
                         Inches(2.1))
    _para(tf, "More on Excel", 60, DARK_BLUE, bold=True, font=TITLE_FONT,
          first=True, space_before=0, space_after=0, line_spacing=1.0)
    _para(tf, "Smart Classroom \u2022 Interactive Learning", 20, BLUE,
          italic=True, space_before=8)

    # Class VI band
    band = add_round_rect(slide, Inches(0.9), Inches(4.85), Inches(3.4),
                          Inches(0.85), GREEN, radius=0.4, shadow=True)
    tb, tf = add_textbox(slide, Inches(0.9), Inches(4.85), Inches(3.4),
                         Inches(0.85), anchor=MSO_ANCHOR.MIDDLE)
    _para(tf, "\U0001F393  Class VI", 24, WHITE, bold=True,
          align=PP_ALIGN.CENTER, first=True, space_before=0, space_after=0)

    # subject label
    tb, tf = add_textbox(slide, Inches(0.9), Inches(5.95), Inches(6.5),
                         Inches(0.6))
    _para(tf, "Computer Science  \u2022  Microsoft Excel", 16, GREY,
          first=True, space_before=0)

    add_footer(slide)


# ---- Learning Activity slide ---------------------------------------------
def build_learning_activity():
    slide = new_slide()
    set_gradient_bg(slide)
    content_header(slide, "\U0001F3AF", "Learning Activity", accent=GREEN)

    # Big central callout
    box = add_round_rect(slide, Inches(2.15), Inches(2.1), Inches(9.0),
                         Inches(2.15), WHITE, line_color=GREEN_LINE, line_w=2,
                         radius=0.12)
    tb, tf = add_textbox(slide, Inches(2.35), Inches(2.1), Inches(8.6),
                         Inches(2.15), anchor=MSO_ANCHOR.MIDDLE)
    _para(tf, "\U0001F4DD  Questions & Answers", 40, DARK_BLUE, bold=True,
          font=TITLE_FONT, align=PP_ALIGN.CENTER, first=True,
          space_before=0, space_after=6)
    _para(tf, "Let's revise Chapter 3 with fun, interactive activities!", 20,
          GREEN, italic=True, align=PP_ALIGN.CENTER, space_before=0)

    # three feature chips
    chips = [
        ("\U0001F4CA", "Excel Skills", BLUE),
        ("\U0001F4DD", "Practice Q&A", ORANGE),
        ("\u2705", "Check Answers", GREEN),
    ]
    chip_w = Inches(3.5)
    gap = Inches(0.35)
    total = chip_w * 3 + gap * 2
    start_x = (SW - total) / 2
    for i, (emo, txt, col) in enumerate(chips):
        cx = start_x + i * (chip_w + gap)
        chip = add_round_rect(slide, cx, Inches(4.7), chip_w, Inches(1.35),
                              WHITE, line_color=col, line_w=2, radius=0.2)
        tb, tf = add_textbox(slide, cx, Inches(4.7), chip_w, Inches(1.35),
                             anchor=MSO_ANCHOR.MIDDLE)
        _para(tf, emo, 30, col, align=PP_ALIGN.CENTER, first=True,
              space_before=0, space_after=2)
        _para(tf, txt, 18, DARK_BLUE, bold=True, align=PP_ALIGN.CENTER,
              space_before=0)

    add_footer(slide)


# ---- Section divider slide -----------------------------------------------
def build_divider(icon, title, subtitle, emojis, accent=BLUE, accent2=GREEN):
    slide = new_slide()
    # bold gradient background
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    rect.shadow.inherit = False
    _no_line(rect)
    rect.fill.gradient()
    st = rect.fill.gradient_stops
    st[0].position = 0.0
    st[0].color.rgb = accent
    st[1].position = 1.0
    st[1].color.rgb = accent2
    try:
        rect.fill.gradient_angle = 45.0
    except Exception:
        pass

    # scattered translucent icons
    positions = [
        (Inches(0.7), Inches(0.6), 40), (Inches(11.7), Inches(0.9), 44),
        (Inches(1.4), Inches(5.7), 40), (Inches(11.2), Inches(5.6), 40),
        (Inches(6.2), Inches(0.5), 34), (Inches(10.0), Inches(3.2), 34),
        (Inches(2.6), Inches(3.0), 30),
    ]
    for (px, py, sz), emo in zip(positions, (emojis * 3)):
        deco_icon(slide, emo, px, py, size=sz,
                  color=RGBColor(0xFF, 0xFF, 0xFF))

    # central white rounded card
    card = add_round_rect(slide, Inches(2.4), Inches(2.15), Inches(8.5),
                          Inches(3.2), WHITE, radius=0.1)
    tb, tf = add_textbox(slide, Inches(2.7), Inches(2.15), Inches(7.9),
                         Inches(3.2), anchor=MSO_ANCHOR.MIDDLE)
    _para(tf, icon, 60, accent, align=PP_ALIGN.CENTER, first=True,
          space_before=0, space_after=2)
    _para(tf, title, 44, DARK_BLUE, bold=True, font=TITLE_FONT,
          align=PP_ALIGN.CENTER, space_before=0, space_after=6)
    # underline accent
    _para(tf, subtitle, 20, accent2, bold=True, italic=True,
          align=PP_ALIGN.CENTER, space_before=0)

    add_footer(slide, dark=True)


# ---- Building blocks for Q&A content -------------------------------------
def question_block(slide, x, y, w, number, qtext, qsize=20):
    """A dark-blue bold question line inside a light box. Returns bottom y."""
    # measure-ish: allocate height by rough char count
    approx_chars_per_line = max(1, int((w / Inches(1)) * 10.2 / (qsize / 20.0)))
    lines = max(1, -(-len(number + "  " + qtext) // approx_chars_per_line))
    h = Inches(0.45) + Inches(0.34) * (lines - 1)
    box = add_round_rect(slide, x, y, w, h, LIGHT_BLUE, radius=0.16,
                         shadow=True)
    tb, tf = add_textbox(slide, x + Inches(0.15), y, w - Inches(0.3), h,
                         anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    p.line_spacing = 1.0
    r = p.add_run()
    r.text = (number + "  ") if number else ""
    r.font.size = Pt(qsize)
    r.font.bold = True
    r.font.name = BODY_FONT
    r.font.color.rgb = ORANGE
    r2 = p.add_run()
    r2.text = qtext
    r2.font.size = Pt(qsize)
    r2.font.bold = True
    r2.font.name = BODY_FONT
    r2.font.color.rgb = DARK_BLUE
    return y + h, box


def options_block(slide, x, y, w, options, size=17):
    """Bulleted MCQ options. options = list of (label, text). Returns bottom y."""
    h = Inches(0.36) * len(options) + Inches(0.08)
    tb, tf = add_textbox(slide, x + Inches(0.35), y, w - Inches(0.5), h)
    first = True
    for label, txt in options:
        p, r = _para(tf, f"\u25CB  {label} {txt}", size, TEXT_DARK,
                     first=first, space_before=1, space_after=1,
                     line_spacing=1.0)
        first = False
    return y + h, tb


def answer_box(slide, x, y, w, answer_text, prefix="\u2705 Answer:", size=19,
               fill=GREEN_BG, line=GREEN_LINE, txt_color=GREEN, tall=None):
    approx_chars_per_line = max(1, int((w / Inches(1)) * 10.6 / (size / 19.0)))
    full = f"{prefix} {answer_text}"
    lines = max(1, -(-len(full) // approx_chars_per_line))
    h = tall if tall else (Inches(0.5) + Inches(0.32) * (lines - 1))
    box = add_round_rect(slide, x, y, w, h, fill, line_color=line, line_w=1.75,
                         radius=0.18, shadow=True)
    tb, tf = add_textbox(slide, x + Inches(0.2), y, w - Inches(0.4), h,
                         anchor=MSO_ANCHOR.MIDDLE)
    p = tf.paragraphs[0]
    p.line_spacing = 1.02
    r = p.add_run()
    r.text = prefix + " "
    r.font.size = Pt(size)
    r.font.bold = True
    r.font.name = BODY_FONT
    r.font.color.rgb = txt_color
    r2 = p.add_run()
    r2.text = answer_text
    r2.font.size = Pt(size)
    r2.font.bold = True
    r2.font.name = BODY_FONT
    r2.font.color.rgb = txt_color
    return y + h, box


def tf_answer_box(slide, x, y, w, is_true, answer_text, size=19):
    """True/False answer box with icon."""
    icon = "\u2705" if is_true else "\u274C"
    fill = GREEN_BG if is_true else RGBColor(0xFD, 0xEA, 0xEA)
    line = GREEN_LINE if is_true else RGBColor(0xE5, 0x73, 0x73)
    col = GREEN if is_true else RED
    return answer_box(slide, x, y, w, answer_text, prefix=f"{icon} Answer:",
                      size=size, fill=fill, line=line, txt_color=col)


# ---------------------------------------------------------------------------
# Content assembly
# ---------------------------------------------------------------------------
def animate(slide, shapes):
    ids = [s.shape_id for s in shapes if s is not None]
    add_fade_animations(slide, ids)


LEFT = Inches(0.6)
FULLW = SW - Inches(1.2)


def mcq_slide(header_icon, header_text, header_sub, items, accent=ORANGE):
    """items: list of dicts {number, q, options:[(label,text)], answer}"""
    slide = new_slide()
    set_gradient_bg(slide)
    content_header(slide, header_icon, header_text, header_sub, accent=accent)
    anim = []
    y = Inches(1.58)
    for it in items:
        y, qbox = question_block(slide, LEFT, y, FULLW, it["number"], it["q"],
                                 qsize=19)
        y += Inches(0.05)
        y, obox = options_block(slide, LEFT, y, FULLW, it["options"], size=17)
        y += Inches(0.05)
        y, abox = answer_box(slide, LEFT + Inches(0.35), y,
                             FULLW - Inches(0.35), it["answer"], size=18)
        y += Inches(0.2)
        anim += [qbox, obox, abox]
    add_footer(slide)
    animate(slide, anim)
    return slide


def tf_slide(header_icon, header_text, header_sub, intro, items, accent=BLUE):
    """items: list of dicts {number, q, is_true, answer}"""
    slide = new_slide()
    set_gradient_bg(slide)
    content_header(slide, header_icon, header_text, header_sub, accent=accent)
    anim = []
    y = Inches(1.6)
    if intro:
        tb, tf = add_textbox(slide, LEFT, y, FULLW, Inches(0.4))
        _para(tf, intro, 17, GREY, italic=True, bold=True, first=True,
              space_before=0, space_after=0)
        y += Inches(0.5)
    for it in items:
        y, qbox = question_block(slide, LEFT, y, FULLW, it["number"], it["q"],
                                 qsize=20)
        y += Inches(0.08)
        y, abox = tf_answer_box(slide, LEFT + Inches(0.35), y,
                                FULLW - Inches(0.35), it["is_true"],
                                it["answer"])
        y += Inches(0.3)
        anim += [qbox, abox]
    add_footer(slide)
    animate(slide, anim)
    return slide


def fill_slide(header_icon, header_text, header_sub, hints, items,
               accent=GREEN):
    slide = new_slide()
    set_gradient_bg(slide)
    content_header(slide, header_icon, header_text, header_sub, accent=accent)
    anim = []
    y = Inches(1.6)
    if hints:
        hb = add_round_rect(slide, LEFT, y, FULLW, Inches(0.62), ORANGE_BG,
                            line_color=ORANGE, line_w=1.5, radius=0.2)
        tb, tf = add_textbox(slide, LEFT + Inches(0.2), y, FULLW - Inches(0.4),
                             Inches(0.62), anchor=MSO_ANCHOR.MIDDLE)
        p = tf.paragraphs[0]
        r = p.add_run()
        r.text = "\U0001F4A1 Hints:  "
        r.font.size = Pt(17)
        r.font.bold = True
        r.font.name = BODY_FONT
        r.font.color.rgb = ORANGE
        r2 = p.add_run()
        r2.text = hints
        r2.font.size = Pt(17)
        r2.font.bold = True
        r2.font.name = BODY_FONT
        r2.font.color.rgb = DARK_BLUE
        anim.append(hb)
        y += Inches(0.8)
    for it in items:
        y, qbox = question_block(slide, LEFT, y, FULLW, it["number"], it["q"],
                                 qsize=19)
        y += Inches(0.08)
        y, abox = answer_box(slide, LEFT + Inches(0.35), y,
                             FULLW - Inches(0.35), it["answer"])
        y += Inches(0.28)
        anim += [qbox, abox]
    add_footer(slide)
    animate(slide, anim)
    return slide


def _bulleted_answer(slide, x, y, w, prefix, lead, bullets, size=18):
    """Answer box containing a lead line + bullet steps. Returns bottom y."""
    n_bul = len(bullets)
    # estimate lines
    cpl = max(1, int((w / Inches(1)) * 10.4 / (size / 18.0)))
    lead_lines = max(1, -(-len(prefix + " " + lead) // cpl)) if lead else 0
    bul_lines = 0
    for b in bullets:
        bul_lines += max(1, -(-len(b) // cpl))
    total_lines = lead_lines + bul_lines
    h = Inches(0.35) + Inches(0.34) * total_lines + Inches(0.12) * n_bul
    box = add_round_rect(slide, x, y, w, h, GREEN_BG, line_color=GREEN_LINE,
                         line_w=1.75, radius=0.1, shadow=True)
    tb, tf = add_textbox(slide, x + Inches(0.25), y + Inches(0.05),
                         w - Inches(0.5), h - Inches(0.1),
                         anchor=MSO_ANCHOR.TOP)
    first = True
    if lead:
        p = tf.paragraphs[0]
        p.line_spacing = 1.03
        p.space_after = Pt(4)
        r = p.add_run()
        r.text = prefix + " "
        r.font.size = Pt(size)
        r.font.bold = True
        r.font.name = BODY_FONT
        r.font.color.rgb = GREEN_DARK
        r2 = p.add_run()
        r2.text = lead
        r2.font.size = Pt(size)
        r2.font.bold = True
        r2.font.name = BODY_FONT
        r2.font.color.rgb = GREEN
        first = False
    for b in bullets:
        p, r = _para(tf, f"\u2714  {b}", size, GREEN, bold=False,
                     first=first, space_before=2, space_after=2,
                     line_spacing=1.03)
        first = False
    return y + h, box


def qa_slide(header_icon, header_text, header_sub, number, qtext, answer=None,
             lead=None, bullets=None, accent=BLUE, qsize=22, asize=20):
    """A single Q with a (possibly long / bulleted) answer."""
    slide = new_slide()
    set_gradient_bg(slide)
    content_header(slide, header_icon, header_text, header_sub, accent=accent)
    anim = []
    y = Inches(1.7)
    y, qbox = question_block(slide, LEFT, y, FULLW, number, qtext, qsize=qsize)
    y += Inches(0.18)
    anim.append(qbox)
    if bullets is not None:
        y, abox = _bulleted_answer(slide, LEFT + Inches(0.25), y,
                                   FULLW - Inches(0.25),
                                   "\u2705 Answer:", lead, bullets, size=asize)
    else:
        y, abox = answer_box(slide, LEFT + Inches(0.25), y,
                             FULLW - Inches(0.25), answer, size=asize)
    anim.append(abox)
    add_footer(slide)
    animate(slide, anim)
    return slide


def build_closing():
    slide = new_slide()
    rect = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, SW, SH)
    rect.shadow.inherit = False
    _no_line(rect)
    rect.fill.gradient()
    st = rect.fill.gradient_stops
    st[0].position = 0.0
    st[0].color.rgb = GREEN
    st[1].position = 1.0
    st[1].color.rgb = BLUE
    try:
        rect.fill.gradient_angle = 45.0
    except Exception:
        pass
    for (px, py, sz), emo in zip(
        [(Inches(1.0), Inches(0.8), 40), (Inches(11.4), Inches(1.0), 44),
         (Inches(1.6), Inches(5.6), 40), (Inches(11.0), Inches(5.6), 42),
         (Inches(6.3), Inches(0.5), 34)],
        ["\U0001F4CA", "\U0001F4C8", "\U0001F9EE", "\U0001F4DD", "\U0001F393"]):
        deco_icon(slide, emo, px, py, size=sz)
    card = add_round_rect(slide, Inches(2.9), Inches(2.3), Inches(7.5),
                          Inches(2.9), WHITE, radius=0.12)
    tb, tf = add_textbox(slide, Inches(3.1), Inches(2.3), Inches(7.1),
                         Inches(2.9), anchor=MSO_ANCHOR.MIDDLE)
    _para(tf, "\U0001F31F  Great Work!", 46, DARK_BLUE, bold=True,
          font=TITLE_FONT, align=PP_ALIGN.CENTER, first=True,
          space_before=0, space_after=6)
    _para(tf, "You have completed Chapter 3 \u2013 More on Excel", 22, GREEN,
          bold=True, align=PP_ALIGN.CENTER, space_before=0, space_after=4)
    _para(tf, "Keep practising your spreadsheet skills!", 18, BLUE,
          italic=True, align=PP_ALIGN.CENTER, space_before=0)
    add_footer(slide, dark=True)


# ---------------------------------------------------------------------------
# MAIN
# ---------------------------------------------------------------------------
def main():
    build_title()
    build_learning_activity()

    # ===== SECTION 1: WARM UP =====
    build_divider("\U0001F4DD", "Warm Up!", "True or False  \u2022  Page 40",
                  ["\U0001F4DD", "\U0001F4CA", "\u2705", "\u274C",
                   "\U0001F4DA"], accent=ORANGE, accent2=RGBColor(0xF2, 0xA6, 0x4B))

    tf_slide("\U0001F4DD", "Warm Up! \u2014 True or False", "Page 40",
             "Write 'T' for true and 'F' for false:",
             [
                 {"number": "1.", "q": "Excel makes calculation easier.",
                  "is_true": True, "answer": "T (True)"},
                 {"number": "2.", "q": "We can apply all types of mathematical operations like addition, subtraction, multiplication and division in Excel.",
                  "is_true": True, "answer": "T (True)"},
                 {"number": "3.", "q": "The input of =25+B2 will give an error message.",
                  "is_true": False,
                  "answer": "F (False) (This is a correct formula that adds 25 to whatever number is inside cell B2)."},
             ], accent=ORANGE)

    # ===== SECTION 2: RAPID FIRE =====
    build_divider("\U0001F525", "Rapid Fire", "Pages 41 and 42",
                  ["\U0001F525", "\U0001F4CA", "\U0001F4DD", "\u2705",
                   "\U0001F9EE"], accent=BLUE, accent2=RGBColor(0x3E, 0x9B, 0xD6))

    # Rapid Fire Q1 - Tick the correct option
    mcq_slide("\U0001F525", "Rapid Fire \u2014 Tick (\u2713) the correct option",
              "1. Tick (\u2713) the correct option \u2022 Page 41",
              [
                  {"number": "a.", "q": "Home tab includes formatting elements such as _____________.",
                   "options": [("(i)", "Text wrap"), ("(ii)", "Merging cells"),
                               ("(iii)", "Cell style"), ("(iv)", "All of the above")],
                   "answer": "(iv) All of the above"},
                  {"number": "b.", "q": "By right-clicking on the sheet tab, we can rename the ____________.",
                   "options": [("(i)", "Workbook"), ("(ii)", "Cell"),
                               ("(iii)", "Worksheet"), ("(iv)", "Column")],
                   "answer": "(iii) Worksheet"},
              ])

    mcq_slide("\U0001F525", "Rapid Fire \u2014 Tick (\u2713) the correct option",
              "1. Tick (\u2713) the correct option \u2022 Page 41",
              [
                  {"number": "c.", "q": "You can specify a column width of 0 to __________.",
                   "options": [("(i)", "255"), ("(ii)", "256"),
                               ("(iii)", "409"), ("(iv)", "290")],
                   "answer": "(i) 255"},
                  {"number": "d.", "q": "We can edit the contents of a cell by ___________ on it.",
                   "options": [("(i)", "Right-clicking"), ("(ii)", "Double-clicking"),
                               ("(iii)", "Clicking"), ("(iv)", "We cannot edit cell content")],
                   "answer": "(ii) Double-clicking"},
              ])

    mcq_slide("\U0001F525", "Rapid Fire \u2014 Tick (\u2713) the correct option",
              "1. Tick (\u2713) the correct option \u2022 Page 41",
              [
                  {"number": "e.", "q": "The ____________ command is used for wrapping text in a cell.",
                   "options": [("(i)", "Wrap Text"), ("(ii)", "Text Wrap"),
                               ("(iii)", "Wrap"), ("(iv)", "Wrapping")],
                   "answer": "(i) Wrap Text"},
              ])

    # Rapid Fire Q2 - True/False
    tf_slide("\U0001F525", "Rapid Fire \u2014 True or False",
             "2. Write 'T' for true and 'F' for false \u2022 Page 42",
             "Write 'T' for true and 'F' for false:",
             [
                 {"number": "a.", "q": "We cannot select more than one cell at a time in an Excel worksheet.",
                  "is_true": False,
                  "answer": "F (False) (We can easily select a group or range of cells)."},
                 {"number": "b.", "q": "We can format number in a cell.",
                  "is_true": True, "answer": "T (True)"},
                 {"number": "c.", "q": "We cannot merge two or more cells.",
                  "is_true": False,
                  "answer": "F (False) (We can merge cells using the 'Merge & Center' command)."},
             ], accent=BLUE)

    tf_slide("\U0001F525", "Rapid Fire \u2014 True or False",
             "2. Write 'T' for true and 'F' for false \u2022 Page 42",
             "Write 'T' for true and 'F' for false:",
             [
                 {"number": "d.", "q": "Auto Fill allows user to automatically fill a number or text series.",
                  "is_true": True, "answer": "T (True)"},
                 {"number": "e.", "q": "A cell or a range of cells that you want to use in your calculation is called formula.",
                  "is_true": False,
                  "answer": "F (False) (It is actually called a Reference)."},
             ], accent=BLUE)

    # ===== SECTION 3: EVALUATION TIME =====
    build_divider("\U0001F4DD", "Evaluation Time", "Pages 42 and 43",
                  ["\U0001F4DD", "\U0001F4CB", "\u270F\uFE0F", "\U0001F4DA",
                   "\U0001F4CA"], accent=GREEN, accent2=RGBColor(0x3E, 0xA8, 0x55))

    # Eval Q1 - Fill in the blanks
    fill_slide("\U0001F4DD", "Evaluation Time \u2014 Fill in the blanks",
               "1. Fill in the blanks using the hints given below \u2022 Page 42",
               "Cell Styles, Select All, Insert, Operators, Merge & Center",
               [
                   {"number": "a.", "q": "The ____________ command on the Home tab is used to add a column into a worksheet.",
                    "answer": "Insert"},
                   {"number": "b.", "q": "We can apply a style on the cells by using the ____________ command on the Styles group under the Home tab.",
                    "answer": "Cell Styles"},
                   {"number": "c.", "q": "We can select entire worksheet by clicking on the ____________ button.",
                    "answer": "Select All"},
               ])

    fill_slide("\U0001F4DD", "Evaluation Time \u2014 Fill in the blanks",
               "1. Fill in the blanks using the hints given below \u2022 Page 42",
               None,
               [
                   {"number": "d.", "q": "We can use ____________ option in the Alignment group to combine two or more selected cells into a single cell.",
                    "answer": "Merge & Center"},
                   {"number": "e.", "q": "Symbols that specify the calculations to be performed are called ____________.",
                    "answer": "Operators"},
               ])

    # Eval Q2 - Short answer
    qa_slide("\U0001F4DD", "Evaluation Time \u2014 Short Answer",
             "2. Short Answer Type Questions \u2022 Page 42", "a.",
             "Can we unmerge the merged cells? If yes, write the name of the command used to do so.",
             answer="Yes, we can unmerge cells. The command used is Unmerge Cells (found under the Merge & Center menu in the Alignment group).",
             accent=GREEN, qsize=22, asize=20)

    qa_slide("\U0001F4DD", "Evaluation Time \u2014 Short Answer",
             "2. Short Answer Type Questions \u2022 Page 43", "b.",
             "What is the use of Copy command on the Home tab?",
             answer="The Copy command is used to duplicate data. It copies the content of a cell to a new place while keeping the original data safe in its old place.",
             accent=GREEN, qsize=22, asize=20)

    qa_slide("\U0001F4DD", "Evaluation Time \u2014 Short Answer",
             "2. Short Answer Type Questions \u2022 Page 43", "c.",
             "Which feature of Excel allows us to display multiple lines of text inside a cell?",
             answer="The Wrap Text feature.",
             accent=GREEN, qsize=22, asize=22)

    # Eval Q3 - Long answer (one per slide, bulleted steps)
    qa_slide("\U0001F4DD", "Evaluation Time \u2014 Long Answer",
             "3. Long Answer Type Questions \u2022 Page 43", "a.",
             "Write the steps to wrap the text in a cell.",
             lead="To wrap text inside a cell, follow these simple steps:",
             bullets=[
                 "Click on the cell containing the text you want to wrap.",
                 "Go to the Home tab.",
                 "Click on the Wrap Text command in the Alignment group.",
             ], accent=GREEN, qsize=22, asize=20)

    qa_slide("\U0001F4DD", "Evaluation Time \u2014 Long Answer",
             "3. Long Answer Type Questions \u2022 Page 43", "b.",
             "How will you change the row height and column width?",
             lead="You can change them by following these steps:",
             bullets=[
                 "Select the rows or columns that you want to change.",
                 "Go to the Home tab and click on the Format command in the Cells group.",
                 "From the drop-down menu, choose Row Height or Column Width.",
                 "Type the number you want in the box and click OK.",
             ], accent=GREEN, qsize=22, asize=19)

    qa_slide("\U0001F4DD", "Evaluation Time \u2014 Long Answer",
             "3. Long Answer Type Questions \u2022 Page 43", "c.",
             "How will you apply cell borders?",
             lead="To add a border around your cells:",
             bullets=[
                 "Select the cells where you want to add a border.",
                 "Go to the Home tab and look for the Font group.",
                 "Click the small arrow next to the Borders command.",
                 "Select a border style (like All Borders) from the list.",
             ], accent=GREEN, qsize=22, asize=19)

    # ===== SECTION 4: COMPETENCY-BASED QUESTIONS =====
    build_divider("\U0001F3AF", "Competency-Based Questions",
                  "Application-Based  \u2022  Page 43",
                  ["\U0001F3AF", "\U0001F4CA", "\U0001F9E9", "\U0001F4DD",
                   "\U0001F4A1"], accent=ORANGE, accent2=RGBColor(0xF0, 0x9B, 0x3E))

    qa_slide("\U0001F3AF", "Competency-Based / Application-Based",
             "4. Competency-Based / Application-Based Questions \u2022 Page 43",
             "a.",
             "Sonia is preparing project details in a spreadsheet. Some text in a cell is too long and gets hidden when she types text in the adjacent cell on the right. How can she change this?",
             lead="Sonia can solve this in two easy ways:",
             bullets=[
                 "She can use the Wrap Text feature to show the text in multiple lines inside the same cell.",
                 "She can increase the Column Width so that all the text fits in one line.",
             ], accent=ORANGE, qsize=21, asize=19)

    qa_slide("\U0001F3AF", "Competency-Based / Application-Based",
             "4. Competency-Based / Application-Based Questions \u2022 Page 43",
             "b.",
             "Anaya has prepared marksheet of her class. She realised that she has forgotten to enter marks of one subject. How can she create space in the marksheet to enter those marks?",
             lead=None,
             answer="Anaya can insert a new column for the forgotten subject. To do this, she should select the column next to where she wants to add the marks, click on the Insert command on the Home tab, and choose Insert Sheet Columns.",
             accent=ORANGE, qsize=21, asize=19)

    build_closing()

    out = "Chapter_3_More_on_Excel_Class_VI.pptx"
    prs.save(out)
    print(f"Saved {out} with {len(prs.slides._sldIdLst)} slides.")


if __name__ == "__main__":
    main()
